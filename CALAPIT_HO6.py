import openpyxl as op

def fave():
    workbook = op.Workbook()
    sheet = workbook.active

    sheet["A1"] = "ID"
    sheet["B1"] = "First Name"
    sheet["C1"] = "Last Name"
    sheet["D1"] = "Birth Year"
    sheet["E1"] = "Age"

    print("===== FAVORITE PEOPLE RECORDER =====")

    for i in range (1,4):
        print(f"\nFavorite Person {i}")
        first_name = input("Enter first name: ")
        last_name = input("Enter last name: ")
        birth_year = int(input("Enter the birth year: "))
        age = 2026 - birth_year

        fav_prsn = [i, first_name, last_name, birth_year, age]
        sheet.append(fav_prsn)

    workbook.save("favorite_people.xlsx")
    print("\nFavorite people added successfully!")

    print("\n===== FAVORITE PEOPLE LISTS =====\n")
    wb_read = op.load_workbook("favorite_people.xlsx")
    ws_read = wb_read.active

    for row in ws_read.iter_rows(values_only=True):
        print(row)
    
    input("Please, press Enter to exit the program......")

fave()
