import tkinter as tk
from tkinter import ttk, messagebox
import openpyxl as op
from datetime import datetime

#GUI
window = tk.Tk()
window.title("Book Borrowing System")
window.configure(bg="pink")

#for displaying
def display():
    workbook = op.load_workbook("Calapit_Database.xlsx")
    sheet = workbook.active

    #clear table
    for row in table.get_children():
        table.delete(row)

    for row in sheet.iter_rows(min_row=2,values_only=True):
        table.insert("",tk.END,values=row)

#function for the late fee kineme
def calculate_late_fee():
    try:
        due = ddate_entry.get()
        returned = returned_entry.get()

        if due == "" or returned == "":
            late_entry.delete(0, tk.END)
            late_entry.insert(0, "0")
            return

        due_date = datetime.strptime(due, "%Y/%m/%d")
        returned_date = datetime.strptime(returned, "%Y/%m/%d")

        if returned_date > due_date:
            fee = (returned_date - due_date).days * 10
        else:
            fee = 0

        late_entry.delete(0, tk.END)
        late_entry.insert(0, str(fee))

    except:
        late_entry.delete(0, tk.END)
        late_entry.insert(0, "0")
        
def input_validation():
    book_name = bkname_entry.get().title()
    author = auname_entry.get().title()
    borrower = brname_entry.get().title()
    contact = cno_entry.get()
    borrow_date = bdate_entry.get()
    due_date = ddate_entry.get()
    returned_date = returned_entry.get()
    status = status_entry.get().title()


    if not book_name:
        messagebox.showerror("Error", "Book Name is required!")
        return False

    if not author:
        messagebox.showerror("Error", "Author Name is required!")
        return False

    if not borrower:
        messagebox.showerror("Error", "Borrower Name is required!")
        return False

    if not contact:
        messagebox.showerror("Error", "Contact Number is required!")
        return False

    if not contact.isdigit():
        messagebox.showerror("Error", "Contact Number must contain numbers only!")
        return False

    if len(contact) != 11:
        messagebox.showerror("Error", "Contact number must be exactly 11 digits!")
        return False

    if not borrow_date:
        messagebox.showerror("Error", "Borrow Date is required!")
        return False

    if not due_date:
        messagebox.showerror("Error", "Due Date is required!")
        return False

    try:
        datetime.strptime(borrow_date, "%Y/%m/%d")
        datetime.strptime(due_date, "%Y/%m/%d")

        if returned_date:
            datetime.strptime(returned_date, "%Y/%m/%d")

    except ValueError:
        messagebox.showerror(
            "Date Error",
            "Please use YYYY/MM/DD format."
        )
        return False
    

    if not status:
        messagebox.showerror("Error", "Status is required!")
        return False

    if returned_date:
        if returned_date <= borrow_date:
            messagebox.showerror("Error", "Returned date must be after borrow date")
            return False
        if returned_date < due_date and status == "Late":
            messagebox.showerror("Error", "Status 'Late' is invalid for early return")
            return False
    
    if status not in ["Borrowed", "Returned", "Late"]:
        messagebox.showerror(
            "Error",
            "Status must only be Borrowed or  Returned, or Late!")
        return False

    return True

def auto_populate(event):
    selected = table.focus()
    values = table.item(selected, "values")

    if not selected:
        return

    #clears the nakalagay sa entry box
    bkname_entry.delete(0, tk.END)
    auname_entry.delete(0, tk.END)
    brname_entry.delete(0, tk.END)
    cno_entry.delete(0, tk.END)
    bdate_entry.delete(0, tk.END)
    ddate_entry.delete(0, tk.END)
    status_entry.delete(0, tk.END)
    returned_entry.delete(0, tk.END)
    late_entry.delete(0, tk.END)

    #ito yung real code sa autopopulate
    if values:
        bkname_entry.insert(0, values[1])
        auname_entry.insert(0, values[2])
        brname_entry.insert(0, values[3])
        cno_entry.insert(0, values[4])
        bdate_entry.insert(0, values[5])
        ddate_entry.insert(0, values[6])
        status_entry.insert(0, values[7])
        returned_entry.insert(0, values[8])
        late_entry.insert(0, values[9])

def saving():
    if not input_validation():
        return

    book_name = bkname_entry.get().title()
    author = auname_entry.get().title()
    borrower = brname_entry.get().title()
    contact = cno_entry.get()
    borrow_date = bdate_entry.get()
    due_date = ddate_entry.get()
    status = status_entry.get().title()
    late_fee = late_entry.get()

    workbook = op.load_workbook("Calapit_Database.xlsx")
    sheet = workbook.active

    #this is for the book id
    book_id = sheet.max_row

    sheet.append([
    book_id,
    book_name,
    author,
    borrower,
    contact,
    borrow_date,
    due_date,
    status,
    returned_entry.get(),
    late_entry.get()])
    workbook.save("Calapit_Database.xlsx")

    messagebox.showinfo("Success", "Record added successfully!")

    calculate_late_fee()
    display()

def update():
    selected = table.focus()

    if not selected:
        messagebox.showerror("Error", "Select a record first!")
        return

    if not input_validation():
        return

    values = table.item(selected, "values")
    book_id = values[0]

    book_name = bkname_entry.get().title()
    author = auname_entry.get().title()
    borrower = brname_entry.get().title()
    contact = cno_entry.get()
    borrow_date = bdate_entry.get()
    due_date = ddate_entry.get()
    status = status_entry.get().title()
    returned_date = returned_entry.get()

    calculate_late_fee()
    late_fee = late_entry.get()

    workbook = op.load_workbook("Calapit_Database.xlsx")
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2):
        if str(row[0].value) == str(book_id):
            row[1].value = book_name
            row[2].value = author
            row[3].value = borrower
            row[4].value = contact
            row[5].value = borrow_date
            row[6].value = due_date
            row[7].value = status
            row[8].value = returned_date
            row[9].value = late_fee
            break

    workbook.save("Calapit_Database.xlsx")

    messagebox.showinfo("Success", "Record updated successfully!")
    display()

def delete():
    selected = table.selection()

    if not selected:
        messagebox.showerror("Error", "Select a record first!")
        return

    values = table.item(selected[0], "values")
    book_id = values[0]

    confirm = messagebox.askyesno(
        "Confirm",
        "Are you sure you want to delete this record?"
    )

    if not confirm:
        return

    workbook = op.load_workbook("Calapit_Database.xlsx")
    sheet = workbook.active

    for i, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if str(row[0].value) == str(book_id):
            sheet.delete_rows(i)
            break

    workbook.save("Calapit_Database.xlsx")

    messagebox.showinfo("Success", "Record deleted successfully!")
    display()

# Form Title
title = tk.Label(window, text="Book Borrowing System", font=("Times New Roman", 15, "bold"), bg="pink")
title.grid(row=0, column=0, columnspan=10, pady=10)

# Frame
genframe = tk.Frame(window, bg="#FFE4E1", bd=2, relief="groove")
genframe.grid(
    row=1,
    column=0,
    columnspan=10,
    pady=10)
genframe.grid_anchor("center")

# Book Name Entry
bkname_entry = tk.Entry(genframe, font=("Times New Roman", 12))
bkname_entry.grid(row=2, column=1, padx=10, pady=(10, 0))

bkname_label = tk.Label(
    genframe,
    text="Book Name",
    font=("Times New Roman", 8,"bold"),
    bg="#FFE4E1")
bkname_label.grid(row=3, column=1)

# Book Author Entry
auname_entry = tk.Entry(genframe, font=("Times New Roman", 12))
auname_entry.grid(row=2, column=2, padx=10, pady=(10, 0))
auname_label = tk.Label(
    genframe,
    text="Author Name",
    font=("Times New Roman", 8, "bold"),
    bg="#FFE4E1")
auname_label.grid(row=3, column=2)

# Borrower Name Entry
brname_entry = tk.Entry(genframe, font=("Times New Roman", 12))
brname_entry.grid(row=2, column=3, padx=10, pady=(10, 0))

brname_label = tk.Label(
    genframe,
    text="Borrower Name",
    font=("Times New Roman", 8,"bold"),
    bg="#FFE4E1")
brname_label.grid(row=3, column=3)

# Borrower Contact Number Entry
cno_entry = tk.Entry(genframe, font=("Times New Roman", 12))
cno_entry.grid(row=2, column=4, padx=10, pady=(10, 0))

cno_label = tk.Label(
    genframe,
    text="Contact Number",
    font=("Times New Roman", 8,"bold"),
    bg="#FFE4E1")
cno_label.grid(row=3, column=4)

# Borrow Date Entry
bdate_entry = tk.Entry(genframe, font=("Times New Roman", 12))
bdate_entry.grid(row=4, column=1, padx=10, pady=(10, 0))

bdate_label = tk.Label(
    genframe,
    text="Borrow Date",
    font=("Times New Roman", 8,"bold"),
    bg="#FFE4E1")
bdate_label.grid(row=5, column=1)

bdate2_label = tk.Label(
    genframe,
    text="(YYYY/MM/DD)",
    font=("Times New Roman",8),
    bg="#FFE4E1")
bdate2_label.grid(row=6, column=1)


# Due Date Entry
ddate_entry = tk.Entry(genframe, font=("Times New Roman", 12))
ddate_entry.grid(row=4, column=2, padx=10, pady=(10, 0))

ddate_label = tk.Label(
    genframe,
    text="Due Date",
    font=("Times New Roman", 8,"bold"),
    bg="#FFE4E1")
ddate_label.grid(row=5, column=2)

ddate2_label = tk.Label(
    genframe,
    text="(YYYY/MM/DD)",
    font=("Times New Roman",8),
    bg="#FFE4E1")
ddate2_label.grid(row=6, column=2)

# Status Entry
status_entry = tk.Entry(genframe, font=("Times New Roman", 12))
status_entry.grid(row=4, column=3, padx=10, pady=(10, 0))

status_label = tk.Label(
    genframe,
    text="Status",
    font=("Times New Roman", 8,"bold"),
    bg="#FFE4E1")
status_label.grid(row=5, column=3)

status2_label = tk.Label(
    genframe,
    text="(Borrowed/Returned/Late)",
    font=("Times New Roman",8),
    bg="#FFE4E1")
status2_label.grid(row=6, column=3)


#Returned Date to
returned_entry = tk.Entry(genframe, font=("Times New Roman", 12))
returned_entry.grid(row=4, column=4, padx=10, pady=(10, 0))

returned_label = tk.Label(
    genframe,
    text="Returned Date",
    font=("Times New Roman", 8,"bold"),
    bg="#FFE4E1")
returned_label.grid(row=5, column=4)

returned2_label = tk.Label(
    genframe,
    text="(YYYY/MM/DD)",
    font=("Times New Roman",8),
    bg="#FFE4E1")
returned2_label.grid(row=6, column=4)

#your late fee kineme
late_label = tk.Label(
    genframe,
    text="Late Return Fee",
    font=("Times New Roman", 8, "bold"),
    bg="#FFE4E1")
late_label.grid(row=8, column=0, columnspan=5, pady=(10, 0))

late_entry = tk.Entry(genframe, font=("Times New Roman", 12), justify="center")
late_entry.grid(row=9, column=0, columnspan=5, pady=(0, 10))

# Buttons
submit_btn = tk.Button(window, 
    text="Submit", 
    font=("Times New Roman", 12, "bold"), 
    bg="#228B22",
    command=saving)
submit_btn.grid(row=6, column=3, pady=(10, 20))

update_btn = tk.Button(
    window,
    text="Update",
    font=("Times New Roman", 12, "bold"),
    bg="#4169E1",
    command=update)
update_btn.grid(row=6, column=4, columnspan=2, pady=(10, 20))

delete_btn = tk.Button(window, 
    text="Delete", 
    bg="#DC143C", 
    fg="white",
    font=("Times New Roman", 12, "bold"),
    command=delete)
delete_btn.grid(row=6, column=6,pady=(10, 20))

#tree view
table = ttk.Treeview(
    window,
    columns=("Book ID" , "Book Name", "Author Name", "Borrower Name", "Contact Number", "Borrow Date", "Due Date", "Status", "Returned Date", "Late Return Fee"),
    show="headings")

for headings in ("Book ID" , "Book Name", "Author Name", "Borrower Name", "Contact Number", "Borrow Date", "Due Date", "Status", "Returned Date", "Late Return Fee"):
    table.heading(headings, text=headings)

for col in table["columns"]:
    table.column(col, width=100, anchor="center")

table.grid(row=7, column=0, columnspan=10, padx=10, pady=10)

table.bind("<<TreeviewSelect>>", auto_populate)
display()
window.mainloop()